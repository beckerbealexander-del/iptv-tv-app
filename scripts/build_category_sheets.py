import urllib.request
import json
import os
import sys
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

SERVER_URL = "http://cf.rilox.sbs"
USERNAME = "fb5940d0a3a0"
PASSWORD = "b1d99e5206"
USER_AGENT = "IPTVSmartersPro/1.0.0 (Linux; Android 11; TV)"

HEADERS = {"User-Agent": USER_AGENT}

PROJECT_ROOT = r"C:\Users\becke\.gemini\antigravity\scratch\iptv-tv-app"
XLSX_PATH = os.path.join(PROJECT_ROOT, "live_tv_categories_export.xlsx")

def clean_channel_name(name):
    orig = name.strip()
    
    # If it's a section divider banner like '##### GENERAL HD/4K #####'
    if orig.startswith('#') and orig.endswith('#'):
        core = orig.strip('#').strip()
        return f"--- [TRENNER] {core} ---"

    cleaned = orig

    # 1. Prefix-Bereinigung (DE, PRIME, JOYN, WOW, SKYGO, SKY GO, RU, ADULT, AT, CH etc.)
    prefix_pattern = r'^(?:DE|PRIME|JOYN|WOW|SKYGO|SKY\s*GO|RU|ADULT|AT|CH)\s*[:\|\-]\s*'
    for _ in range(3):
        cleaned = re.sub(prefix_pattern, '', cleaned, flags=re.IGNORECASE)

    # Führende Tags ohne Doppelpunkt entfernen
    for p in ['DE ', 'PRIME ', 'WOW ', 'JOYN ', 'SKYGO ', 'SKY GO ', 'RU ', 'ADULT ']:
        if cleaned.upper().startswith(p):
            cleaned = cleaned[len(p):].strip()

    # Nachgestellte Tags wie '| DE' oder ' - DE' entfernen
    cleaned = re.sub(r'\s*[\|\-]\s*DE\b', '', cleaned, flags=re.IGNORECASE)

    # 2. Entferne Provider-Auflösungs-Tags und Marketing-Anhänge aus dem Namen
    cleaned = re.sub(r'\s*\(\s*(?:LOW\s*BIT|LOWBIT|720[Pp]|1080[Pp]|3840[Pp]|SAT)\s*\)', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'\b(?:4K|UHD|FHD|HD|SD|RAW|HEVC|60FPS|50FPS|720[Pp]|1080[Pp]|3840[Pp])\b', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'[ᴴᴰ⁴ᴷᶠʰᵈˢᵈᴿᴬᵂʰᵉᵛᶜᵁᴴᴰ³⁸⁴⁰ᴾ¹⁰⁸⁰ᴾ⁷²⁰ᴾ⁶⁰ᶠᵖˢ⁵⁰ᶠᵖˢ◉]', '', cleaned)

    # Bereinigung doppelter Leerzeichen und Randzeichen
    cleaned = re.sub(r'\s+', ' ', cleaned).strip(' -|:')
    return cleaned

def sanitize_sheet_name(name):
    clean = re.sub(r'[\\/\?\*\:\[\]]', '-', name)
    return clean[:31].strip()

def main():
    print("1. Lade bestehende Excel-Arbeitsmappe mit deinen Zuordnungen...")
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws_main = wb["Live TV Kategorien"]

    categories_by_main = {}
    for row in ws_main.iter_rows(min_row=2, values_only=True):
        orig_name, cat_id, lang, count, main_cat, sort_order = row[:6]
        if not main_cat or str(main_cat).strip().lower() == 'kann entfernt werden':
            continue
        
        main_cat_clean = str(main_cat).strip()
        sort_val = 999 if sort_order is None else int(sort_order)
        key = (sort_val, main_cat_clean)
        
        categories_by_main.setdefault(key, []).append({
            'name': orig_name,
            'id': str(cat_id),
            'count': count
        })

    sorted_groups = sorted(categories_by_main.keys(), key=lambda x: (x[0], x[1]))

    print("\n2. Lade alle Live-Streams von der Xtream-API...")
    streams_url = f"{SERVER_URL}/player_api.php?username={USERNAME}&password={PASSWORD}&action=get_live_streams"
    req = urllib.request.Request(streams_url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=60) as resp:
        all_streams = json.loads(resp.read().decode('utf-8'))
    print(f"-> {len(all_streams)} Streams empfangen.")

    streams_by_category = {}
    for s in all_streams:
        cid = str(s.get('category_id'))
        streams_by_category.setdefault(cid, []).append(s)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B71C1C", end_color="B71C1C", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D8D8D8'),
        right=Side(style='thin', color='D8D8D8'),
        top=Side(style='thin', color='D8D8D8'),
        bottom=Side(style='thin', color='D8D8D8')
    )

    data_font = Font(name="Calibri", size=10)
    data_font_bold = Font(name="Calibri", size=10, bold=True)
    alt_fill = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    existing_sheets = [s for s in wb.sheetnames if s != "Live TV Kategorien"]
    for s in existing_sheets:
        del wb[s]

    print("\n3. Erstelle Blätter für jede Überkategorie...")
    headers = ["Unterkategorie", "Originaler Sendername", "Bereinigter Name"]

    for sort_val, main_name in sorted_groups:
        sheet_title = sanitize_sheet_name(f"{sort_val}. {main_name}" if sort_val != 999 else main_name)
        ws = wb.create_sheet(title=sheet_title)

        ws.sheet_properties.tabColor = "B71C1C"
        ws.freeze_panes = "A2"

        ws.append(headers)
        ws.row_dimensions[1].height = 26

        for col_idx in range(1, 4):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        subcategories = categories_by_main[(sort_val, main_name)]
        current_row = 2

        for subcat in subcategories:
            subcat_id = subcat['id']
            subcat_name = subcat['name']
            ch_list = streams_by_category.get(subcat_id, [])

            for s in ch_list:
                orig_ch_name = s.get('name', '').strip()
                cleaned_ch_name = clean_channel_name(orig_ch_name)

                row_fill = alt_fill if current_row % 2 == 0 else white_fill

                ws.append([subcat_name, orig_ch_name, cleaned_ch_name])
                ws.row_dimensions[current_row].height = 20

                c1 = ws.cell(row=current_row, column=1)
                c2 = ws.cell(row=current_row, column=2)
                c3 = ws.cell(row=current_row, column=3)

                for cell in [c1, c2, c3]:
                    cell.border = thin_border
                    cell.fill = row_fill
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                c1.font = data_font
                c2.font = data_font
                c3.font = data_font_bold

                current_row += 1

        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 46
        ws.column_dimensions["C"].width = 44

        print(f"-> Blatt '{sheet_title}': {current_row - 2} Sender eingetragen.")

    print("\n4. Speichere aktualisierte Excel-Datei...")
    wb.save(XLSX_PATH)
    print(f"ERFOLGREICH ABGESCHLOSSEN!\nPfad: {XLSX_PATH}")

if __name__ == "__main__":
    main()
