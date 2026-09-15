import urllib.request
import json
import csv
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure stdout uses utf-8
sys.stdout.reconfigure(encoding='utf-8')

SERVER_URL = "http://cf.rilox.sbs"
USERNAME = "fb5940d0a3a0"
PASSWORD = "b1d99e5206"
USER_AGENT = "IPTVSmartersPro/1.0.0 (Linux; Android 11; TV)"

HEADERS = {
    "User-Agent": USER_AGENT
}

PROJECT_ROOT = r"C:\Users\becke\.gemini\antigravity\scratch\iptv-tv-app"
CSV_EXPORT_PATH = os.path.join(PROJECT_ROOT, "live_tv_categories_export.csv")
XLSX_EXPORT_PATH = os.path.join(PROJECT_ROOT, "live_tv_categories_export.xlsx")

def fetch_json(action):
    url = f"{SERVER_URL}/player_api.php?username={USERNAME}&password={PASSWORD}&action={action}"
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode('utf-8'))

def classify_category(name):
    u = name.upper()
    
    # Check for excluded international prefixes first
    excluded_prefixes = ["UK|", "US|", "IT|", "TR|", "PL|", "IE|", "AR|", "FR|", "ES|", "PT|", "NL|", "GR|", "EX-YU|", "AL|", "RO|"]
    if any(u.startswith(p) for p in excluded_prefixes):
        return None

    # Deutschsprachig (DE, AT, CH und typische Muster)
    if (u.startswith("DE|") or u.startswith("DE:") or u.startswith("DE -") or u.startswith("DE ") or
        u.startswith("AT|") or u.startswith("AT:") or u.startswith("AT -") or
        u.startswith("CH|") or u.startswith("CH:") or u.startswith("CH -") or
        "GERMAN" in u or "DEUTSCHLAND" in u or "DEUTSCH" in u):
        return "DE"

    # Russischsprachig (RU und typische Muster)
    if (u.startswith("RU|") or u.startswith("RU:") or u.startswith("RU -") or u.startswith("RU ") or
        "RUSSIAN" in u or "RUSSLAND" in u or "RUS" in u):
        return "RU"

    # For Adults / Erotik
    if ("ADULT" in u or "XXX" in u or "18+" in u or "EROTIC" in u or "PORN" in u):
        return "ADULT"

    return None

def main():
    print("1. Lade alle Live-TV-Kategorien von der Xtream-API...")
    categories = fetch_json("get_live_categories")
    print(f"-> {len(categories)} Gesamtkategorien empfangen.")

    print("2. Lade alle Live-Streams zur Ermittlung der Senderanzahl...")
    streams = fetch_json("get_live_streams")
    print(f"-> {len(streams)} Gesamtsender empfangen.")

    # Zähle Sender pro category_id
    channel_counts = {}
    for s in streams:
        cid = str(s.get("category_id"))
        channel_counts[cid] = channel_counts.get(cid, 0) + 1

    # Filtern & Zuordnen
    print("3. Filtere relevante Kategorien (DE, RU, ADULT)...")
    rows = []
    for c in categories:
        cid = str(c.get("category_id"))
        cname = c.get("category_name", "").strip()
        lang_group = classify_category(cname)
        if lang_group:
            count = channel_counts.get(cid, 0)
            rows.append({
                "Original_Category_Name": cname,
                "Category_ID": cid,
                "Language_Group": lang_group,
                "Channel_Count": count,
                "New_Main_Category": "",
                "Sort_Order": ""
            })

    print(f"-> {len(rows)} relevante Kategorien extrahiert.")

    # 4. CSV generieren (UTF-8-SIG mit BOM für perfekten Excel-Import)
    print("4. Schreibe CSV-Datei...")
    headers = [
        "Original_Category_Name",
        "Category_ID",
        "Language_Group",
        "Channel_Count",
        "New_Main_Category",
        "Sort_Order"
    ]

    with open(CSV_EXPORT_PATH, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    # 5. Excel (.xlsx) generieren mit professionellem Styling
    print("5. Schreibe Excel (.xlsx) Datei...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Live TV Kategorien"

    # Header Styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B71C1C", end_color="B71C1C", fill_type="solid") # Dark Red
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data Rows
    de_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alt_fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
    
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)

    for r_idx, row_data in enumerate(rows, 2):
        ws.row_dimensions[r_idx].height = 22
        current_fill = alt_fill if r_idx % 2 == 0 else de_fill
        
        ws.append([
            row_data["Original_Category_Name"],
            row_data["Category_ID"],
            row_data["Language_Group"],
            row_data["Channel_Count"],
            row_data["New_Main_Category"],
            row_data["Sort_Order"]
        ])

        # Cell Formats
        cell_name = ws.cell(row=r_idx, column=1)
        cell_id = ws.cell(row=r_idx, column=2)
        cell_lang = ws.cell(row=r_idx, column=3)
        cell_count = ws.cell(row=r_idx, column=4)
        cell_new_main = ws.cell(row=r_idx, column=5)
        cell_sort = ws.cell(row=r_idx, column=6)

        for c in [cell_name, cell_id, cell_lang, cell_count, cell_new_main, cell_sort]:
            c.border = thin_border
            c.font = data_font
            c.fill = current_fill

        cell_name.alignment = Alignment(horizontal="left", vertical="center")
        cell_id.alignment = Alignment(horizontal="center", vertical="center")
        cell_lang.alignment = Alignment(horizontal="center", vertical="center")
        cell_lang.font = bold_font
        cell_count.alignment = Alignment(horizontal="right", vertical="center")
        cell_new_main.alignment = Alignment(horizontal="left", vertical="center")
        cell_sort.alignment = Alignment(horizontal="center", vertical="center")

    # Spaltenbreiten anpassen
    column_widths = {
        "A": 46,  # Original_Category_Name
        "B": 14,  # Category_ID
        "C": 18,  # Language_Group
        "D": 16,  # Channel_Count
        "E": 28,  # New_Main_Category
        "F": 14   # Sort_Order
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(XLSX_EXPORT_PATH)

    print("\nERFOLGREICH ABGESCHLOSSEN!")
    print(f"CSV-Datei:   {CSV_EXPORT_PATH}")
    print(f"Excel-Datei: {XLSX_EXPORT_PATH}")
    print(f"Gefilterte Kategorien: {len(rows)}")

if __name__ == "__main__":
    main()
