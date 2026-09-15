import urllib.request
import json
import os
import sys
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

SERVER_URL = "http://cf.rilox.sbs"
USERNAME = "fb5940d0a3a0"
PASSWORD = "b1d99e5206"
USER_AGENT = "IPTVSmartersPro/1.0.0 (Linux; Android 11; TV)"
HEADERS = {"User-Agent": USER_AGENT}

PROJECT_ROOT = r"C:\Users\becke\.gemini\antigravity\scratch\iptv-tv-app"
XLSX_PATH = os.path.join(PROJECT_ROOT, "live_tv_categories_export.xlsx")
M3U_CLEAN_PATH = os.path.join(PROJECT_ROOT, "playlist_clean.m3u8")
M3U_MULTI_PATH = os.path.join(PROJECT_ROOT, "playlist_multistream.m3u8")

# Blacklist: Kategorien, die komplett und dauerhaft entfernt werden sollen
BLACKLIST_CATEGORY_IDS = {
    '570',   # AT| DAZN PPV
    '1760',  # DE| SPOTIFY INF & ᴿᴬᵂ
    '17',    # CH| SWITZERLAND HD/4K
    '1861',  # CH| MYSPORTS ᴿᴬᵂ
    '1714',  # CH| BLUE SPORT ᴿᴬᵂ
    '1862',  # CH| BLUE SPORT DIRECT ᴿᴬᵂ
    '1962',  # CH| SFL PPV
    '972',   # CH| DAZN PPV
    '1631',  # AT| CANAL+ ONLINE UNTERHALTUNG ᴿᴬᵂ
    '1630',  # AT| CANAL+ ONLINE SPORT ᴿᴬᵂ
    '2033',  # AT| JOYN ᴿᴬᵂ
    '1955',  # AT| AUSTRIA ⱽᴵᴾ
    '54',    # AT| AUSTRIA HD/4K
}

def clean_channel_name(name, default_main=None):
    orig = name.strip()
    if orig.startswith('#') and orig.endswith('#'):
        core = orig.strip('#').strip()
        return f"--- [TRENNER] {core} ---"

    orig_u = orig.upper()
    is_ru = (default_main == 'Russian') or orig_u.startswith('RU:') or orig_u.startswith('RU ')

    # Russian: Originale Sendernamen beibehalten
    if is_ru:
        return orig

    cleaned = orig
    prefix_pattern = r'^(?:DE|PRIME|JOYN|WOW|SKYGO|SKY\s*GO|RU|ADULT|AT|CH|UK)\s*[:\|\-]\s*'
    for _ in range(3):
        cleaned = re.sub(prefix_pattern, '', cleaned, flags=re.IGNORECASE)

    for p in ['DE ', 'PRIME ', 'WOW ', 'JOYN ', 'SKYGO ', 'SKY GO ', 'RU ', 'ADULT ', 'UK ']:
        if cleaned.upper().startswith(p):
            cleaned = cleaned[len(p):].strip()

    cleaned = re.sub(r'\s*[\|\-]\s*(?:DE|UK)\b', '', cleaned, flags=re.IGNORECASE)

    # Fake-Auflösungen & Suffixe entfernen (inkl. KABEL, WEB 1080, WEB 720P)
    cleaned = re.sub(r'\s*\(\s*(?:MOBIL|MOBILE|LOW\s*BIT|LOWBIT|LOW|720[Pp]|1080[Pp]|3840[Pp]|SAT|KABEL|WEB\s*1080|WEB\s*720[Pp]|WEB)\s*\)', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'\b(?:4K|UHD|FHD|HD|SD|RAW|HEVC|60FPS|50FPS|720[Pp]|1080[Pp]|3840[Pp])\b', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'[ᴴᴰ⁴ᴷᶠʰᵈˢᵈᴿᴬᵂʰᵉᵛᶜᵁᴴᴰ³⁸⁴⁰ᴾ¹⁰⁸⁰ᴾ⁷²⁰ᴾ⁶⁰ᶠᵖˢ⁵⁰ᶠᵖˢ◉]', '', cleaned)

    cleaned = re.sub(r'\s+', ' ', cleaned).strip(' -|:')
    upper_c = cleaned.upper()

    # Deutsche Sender-Synonyme vereinheitlichen
    if upper_c in ['KABEL 1', 'KABEL EINS']:
        cleaned = 'KABEL EINS'
    elif upper_c in ['KABEL 1 CLASSICS', 'KABEL EINS CLASSICS']:
        cleaned = 'KABEL EINS CLASSICS'
    elif upper_c in ['KABEL 1 DOKU', 'KABEL EINS DOKU']:
        cleaned = 'KABEL EINS DOKU'
    elif upper_c in ['RTL 2', 'RTL ZWEI']:
        cleaned = 'RTL ZWEI'
    elif upper_c in ['RTL NITRO', 'NITRO', 'NTRO']:
        cleaned = 'NITRO'
    elif upper_c in ['SUPER RTL', 'RTL SUPER', 'TOGGO RTL SUPER']:
        cleaned = 'RTL SUPER'
    elif upper_c in ['N-TV', 'NTV']:
        cleaned = 'N-TV'
    elif upper_c in ['VOX UP', 'VOXUP']:
        cleaned = 'VOX UP'
    elif upper_c in ['ZDF INFO', 'ZDFINFO']:
        cleaned = 'ZDFINFO'
    elif upper_c in ['ZDF NEO', 'ZDFNEO']:
        cleaned = 'ZDFNEO'
    elif upper_c in ['E! ENTERTAINM', 'E ENTERTAINM', 'E! ENTERTAINMENT']:
        cleaned = 'E! ENTERTAINMENT'
    elif upper_c in ['GUTE LAUNE', 'GUTE LAUNE TV']:
        cleaned = 'GUTE LAUNE TV'

    return cleaned

def normalize_key(clean_name):
    if not clean_name or clean_name.startswith('---'):
        return None
    k = str(clean_name).upper().strip()
    k = re.sub(r'[^A-Z0-9\+]', '', k)
    return k

# Verifizierte Dokusender
DOKU_CHANNELS = {
    'ZDFINFO': 'ZDFinfo',
    'KABELEINSDOKU': 'Kabel Eins Doku',
    'N24DOKU': 'N24 Doku',
    'WELTDERWUNDER': 'Welt der Wunder',
    'FOCUSTVREPORTAGE': 'Focus TV Reportage',
    'DISCOVERY': 'Discovery Channel',
    'DISCOVERYCHANNEL': 'Discovery Channel',
    'NATIONALGEOGRAPHIC': 'National Geographic',
    'NATGEOWILD': 'Nat Geo Wild',
    'ANIMALPLANET': 'Animal Planet',
    'HISTORY': 'History Channel',
    'HISTORYPLAY': 'History Channel',
    'SKYDOCUMENTARIES': 'Sky Documentaries',
    'SKYNATURE': 'Sky Nature',
    'SKYCRIME': 'Sky Crime',
    'GEOTELEVISION': 'Geo Television',
    'GEO': 'Geo Television',
    'SPIEGELGESCHICHTE': 'Spiegel Geschichte',
    'SPIEGELTVGESCHICHTE': 'Spiegel Geschichte',
    'SPIEGELTVKONFLIKTE': 'Spiegel TV Konflikte',
    'CURIOSITYCHANNELPOWEREDBYSPIEGEL': 'Curiosity Channel',
    'CURIOSITYNOW': 'Curiosity Now',
    'MARCOPOLOTV': 'Marco Polo TV',
    'TRAVELXP': 'Travelxp 4K',
    'TRAVELXP4K': 'Travelxp 4K',
    'BERGBLICK': 'Bergblick',
    'TERRAMATERWILD': 'Terra Mater Wild',
    'BBCHISTORY': 'BBC History',
    'BBCTRAVEL': 'BBC Travel',
    'MAGELLANTVNOW': 'MagellanTV Now',
    'ONETERRA': 'One Terra',
    'XPLORE': 'Xplore',
    'LOVETHEPLANET': 'Love the Planet',
    'INSIGHT': 'Insight TV',
    'INSIGHTTV': 'Insight TV',
    'WEDOTVBIGSTORIES': 'wedotv Big Stories',
    'WAIDWERK': 'Waidwerk',
    'CRIME+INVESTIGATION': 'Crime & Investigation',
    'CRIME+INVESTIGATIONPLAY': 'Crime & Investigation',
    'CRIMEINVESTIGATION': 'Crime & Investigation',
    'CRIMEINVESTIGATIONPLAY': 'Crime & Investigation',
    'CRIMESCENETV': 'Crime Scene TV',
    'TOPTRUECRIME': 'Top True Crime',
    'AMERICANCRIMES': 'American Crimes',
    'FILMRISETRUECRIME': 'FilmRise True Crime',
    'TAETERJAGDCRIMESCENESOLVERS': 'Täterjagd (Crime Scene Solvers)',
    'TTERJAGDCRIMESCENESOLVERS': 'Täterjagd (Crime Scene Solvers)',
}

SPECIAL_REROUTES = {
    'SYFY': 'Sky',
    'TLC': 'FreeTV',
}

FREETV_ROUTING_KEYS = {
    'PROSIEBEN', 'PROSIEBENMAXX', 'RTL', 'RTLZWEI', 'SAT1GOLD', 'SAT1', 'VOX',
    'KABELEINS', 'SIXX', 'DMAX', 'COMEDYCENTRAL', 'EENTERTAINM', 'EENTERTAINMENT',
    'NTV', 'NITRO', 'WELT', 'SUPERRTL', 'RTLSUPER', 'TELE5', 'TLC', 'ATV'
}

KIDS_ROUTING_KEYS = {'KIKA', 'DISNEYCHANNEL'}

SPORT_ROUTING_KEYS = {
    'SPORT1', 'EUROSPORT1', 'EUROSPORT2', 'DAZNBAR1', 'DAZNBAR2', 'SKYSPORTNEWS',
    'SKYSPORTPREMIERLEAGUE', 'SKYSPORTTOPEVENT', 'SKYSPORTBUNDESLIGA', 'SKYSPORTF1',
    'SKYSPORTGOLF', 'SKYSPORTMIX', 'SKYSPORTTENNIS', 'SKYSPORTAUSTRIA1',
    'SKYSPORTAUSTRIA2', 'SKYSPORTAUSTRIA3', 'SKYSPORTAUSTRIA4'
}
for i in range(1, 11):
    SPORT_ROUTING_KEYS.add(f'SKYSPORTBUNDESLIGA{i}')
    SPORT_ROUTING_KEYS.add(f'SKYSPORT{i}')

GERMAN_MUSIC_KEYS = {
    'DELUXEMUSIC', 'MTV', 'JUKEBOX', 'GOLDSTARTV', 'GUTELAUNE', 'GUTELAUNETV', 'STINGRAYCLASSICA'
}

def evaluate_source(orig_name, subcat_name):
    orig = str(orig_name)
    sub = str(subcat_name)
    orig_u = orig.upper()
    sub_u = sub.upper()
    
    if 'LOW BIT' in orig_u or 'LOWBIT' in orig_u or '(LOW' in orig_u:
        return (20, 'DVB Satellit (Low Bit / Notfall)')
    if 'MOBIL' in orig_u or 'MOBILE' in orig_u:
        return (30, 'Mobil Stream (Niedrige Bandbreite)')
    if re.search(r'\bSD\b', orig_u) or 'ˢᵈ' in orig:
        return (40, 'SD Feed (Notfall / Geringe Bandbreite)')
    if 'WEB 720' in orig_u:
        return (55, 'Web Direktfeed (720p HD)')
    if '720P' in orig_u or '(720' in orig_u or '⁷²⁰ᴾ' in orig:
        return (50, 'DVB Satellit (720p Feed)')
    if 'PRIME' in orig_u or 'PRIME' in sub_u:
        return (100, 'Prime Video (60fps RAW / Top-Qualität)')
    if 'WOW' in orig_u or 'WOW' in sub_u:
        return (95, 'WOW Direktstream (Dolby Audio)')
    if '4K' in orig_u or '⁴ᴷ' in orig or 'UHD' in orig_u or 'ᵁᴴᴰ' in orig or '3840' in orig:
        if 'SKYGO' in orig_u or 'SKY GO' in orig_u or 'SKYGO' in sub_u:
            return (92, 'SkyGo Web (4K Feed)')
        return (90, '4K / UHD Feed')
    if 'WEB 1080' in orig_u:
        return (89, 'Web Direktfeed (1080p RAW)')
    if 'DAZN EXCLUSIVE' in sub_u or 'DAZN EXKLUSIV' in sub_u or 'DAZN EXCLUSIVE' in orig_u:
        return (88, 'DAZN Exklusiv (RAW Feed)')
    if 'RTL+' in orig_u or 'RTL+' in sub_u:
        return (86, 'RTL+ Premium (RAW Gold)')
    if 'JOYN' in orig_u or 'JOYN' in sub_u:
        return (85, 'Joyn OTT (RAW / Sehr stabil)')
    if '(KABEL' in orig_u:
        return (83, 'DVB Kabel (RAW Direct)')
    if 'RAW' in orig_u or 'ᴿᴬᵂ' in orig:
        return (82, 'RAW Webfeed')
    if 'SKYGO' in orig_u or 'SKY GO' in orig_u or 'SKYGO' in sub_u:
        return (80, 'SkyGo Web (HD Feed)')
    if 'SAT' in orig_u or '(SAT)' in orig_u:
        return (72, 'DVB Satellit (Standard HD)')
    if 'HEVC' in orig_u or 'ʰᵉᵛᶜ' in orig or 'HEVC' in sub_u:
        return (65, 'HEVC H.265 (Kompakt / Sparsam)')
    if 'HD' in orig_u or 'ᴴᴰ' in orig:
        return (75, 'Standard HD Feed')
    return (60, 'Standard Stream')

def auto_classify_category(name):
    u = name.upper().strip()
    is_de = any(u.startswith(p) for p in ['DE|', 'DE:', 'DE -', 'DE ', 'AT|', 'AT:', 'CH|', 'GERMAN', 'DEUTSCH'])
    is_ru = any(u.startswith(p) for p in ['RU|', 'RU:', 'RU -', 'RU ', 'RUSSIAN', 'RUSSLAND'])
    is_adult = any(w in u for w in ['ADULT', 'XXX', '18+', 'PORN', 'EROTIC']) and not any(w in u for w in ['ADULT SWIM', 'ADULT-SWIM'])
    
    if not (is_de or is_ru or is_adult or 'UK| MUSIC' in u):
        return None, None
    if any(w in u for w in ['SPOTIFY', 'DAZN PPV', 'SWITZERLAND', 'BLUE SPORT', 'MYSPORTS', 'SFL PPV', 'CANAL+ ONLINE']):
        return None, None

    if is_ru: return 'Russian', 2
    if is_adult: return 'Privat', 10
    if '24/7' in u: return '24/7 Filme&Serien', 9
    if any(w in u for w in ['SPORT', 'BUNDESLIGA', 'DAZN', 'FUSSBALL', 'TENNIS', 'GOLF']): return 'Sport', 3
    if any(w in u for w in ['DOKU', 'DOC', 'DOCUMENTARY', 'WISSEN']): return 'Doku', 8
    if any(w in u for w in ['PPV', 'LEAGUES', 'EVENT', 'SOCCER PPV']): return 'Live Events', 7
    if any(w in u for w in ['MUSIK', 'MUSIC', 'RADIO', 'SOUND']): return 'Musik', 6
    if any(w in u for w in ['KIDS', 'KINDER', 'CARTOON', 'DISNEY', 'TOGGO']): return 'Kids', 5
    if any(w in u for w in ['SKY', 'CINEMA', 'MOVIES', 'SERIEN', 'WOW ENTERTAINMENT']): return 'Sky', 4
    return 'FreeTV', 1

def get_channel_sort_priority(clean_name, category_key):
    u = clean_name.upper().strip()
    norm_k = re.sub(r'[^A-Z0-9\+]', '', u)

    if 'FREETV' in category_key:
        lcn_top = [
            'DASERSTE', 'ZDF', 'RTL', 'SAT1', 'PROSIEBEN', 'VOX', 'KABELEINS', 'RTLZWEI',
            '3SAT', 'ARTE', 'NITRO', 'DMAX', 'SIXX', 'SAT1GOLD', 'PROSIEBENMAXX', 'VOXUP',
            'RTLUP', 'TELE5', 'SERVUSTV', 'DF1', 'PROSIEBENFUN', 'SAT1EMOTIONS',
            'KABELEINSCLASSICS', 'NTV', 'WELT', 'PHOENIX', 'TAGESSCHAU24', 'ARDALPHA',
            'EURONEWS', 'ZDFNEO', 'ONE', 'KIKA', 'RTLSUPER', 'TOGGOPLUS', 'DISNEYCHANNEL', 'NICKELODEON',
            'TLC', 'RIC', 'EUROSPORT1', 'SPORT1', 'REDBULLTV', 'MTV', 'DELUXEMUSIC'
        ]
        if norm_k in lcn_top:
            return (1, lcn_top.index(norm_k), clean_name)

        rtl_plus = [
            'RTLCRIME', 'RTLLIVING', 'RTLPASSION',
            'RTLCOMEDY', 'RTLHAUSGARTEN', 'RTLSHINE', 'NOW',
            'BAUERSUCHTFRAU', 'ALARMFUERCOBRA11BALKO', 'ALARMFUERCOBRA11',
            'ALLESWASZAEHLTCLASSICS', 'HUNDKATZEMAUS', 'SHOPPINGQUEEN'
        ]
        for idx, rk in enumerate(rtl_plus):
            if rk in norm_k or norm_k in rk:
                return (2, idx, clean_name)

        dritte = ['WDR', 'NDR', 'BRFERNSEHEN', 'BR', 'SWR', 'MDR', 'HR', 'RBB', 'RADIOBREMEN', 'SR']
        for idx, d in enumerate(dritte):
            if d in norm_k:
                return (3, idx, clean_name)

        if u.startswith('SAT.1 ') or u.startswith('RTL '):
            return (4, 0, clean_name)

        lokal = ['ALLGAEU', 'AUGSBURG', 'MUENCHEN', 'FRANKEN', 'RHEIN', 'RNF', 'REGIONAL', 'KTV', '25KTV']
        for idx, lk in enumerate(lokal):
            if lk in norm_k:
                return (5, idx, clean_name)

        shopping = ['QVC', 'HSE', 'SONNENKLAR', 'BIBEL', 'SHOP']
        for idx, sh in enumerate(shopping):
            if sh in norm_k:
                return (6, idx, clean_name)

        return (7, 0, clean_name)

    elif 'RUSSIAN' in category_key:
        u_norm = re.sub(r'^(?:RU|DE)[:\|\-\s]*', '', u, flags=re.IGNORECASE)
        u_norm = re.sub(r'\b(HD|4K|SD|FHD|UHD)\b', '', u_norm, flags=re.IGNORECASE).strip()
        ru_k = re.sub(r'[^A-Z0-9\+]', '', u_norm)

        federal_ru = [
            'PERVIYKANAL', 'PERVIY', 'ROSSIYA1', 'ROSSIYA', 'RTRPLANETA', 'RTR',
            'HTB', 'NTV', 'NTVMIR', 'THT', 'TNT', 'THT4', 'TNT4', 'CTC', 'STS', 'CTSLOVE', 'STSLOVE',
            'PEHTB', 'RENTV', '5KANAL', 'PYATITKANAL', 'PYATIYKANAL', 'MATCHTV', 'MATCH',
            'ROSSIYA24', 'RUSSIAK', 'KULTURA', 'ZVEZDA', 'TV3', 'PYATNITSA', 'DOMASHNIY',
            'CHE', 'KARUSEL', 'MIR', 'MIR24', 'TVCI', 'RBKTV', 'RBK', '2X2', 'UTV', 'U', 'SUPER'
        ]
        for idx, f in enumerate(federal_ru):
            if ru_k == f or ru_k.startswith(f):
                return (1, idx, clean_name)

        sport_ru = ['MATCH', 'KHL', 'SETANTA', 'FUTBOL', 'FOOTBALL', 'SPORT', 'BOKS', 'BOX', 'BOEC', 'ARENA', 'IGRA', 'EUROSPORT', 'EXTRIM']
        for idx, s in enumerate(sport_ru):
            if s in ru_k:
                return (2, idx, clean_name)

        cinema_ru = ['AMEDIA', 'DOMKINO', 'TV1000', 'VIP', 'KINO', 'FOX', 'PARAMOUNT', 'SONY', 'SERIAL', 'ILLUZION', 'AMC', 'A1', 'A2', 'RUSSK']
        for idx, c in enumerate(cinema_ru):
            if c in ru_k:
                return (3, idx, clean_name)

        doku_ru = ['DISCOVERY', 'NATGEO', 'NATIONALGEOGRAPHIC', 'VIASAT', 'NAUKA', 'HISTORY', 'ANIMAL', 'DOC', 'TRAVEL', 'GEOGRAPHIC', 'EUREKA', 'PRIKLYUCHENIYA', 'GLAZAMI']
        for idx, d in enumerate(doku_ru):
            if d in ru_k:
                return (4, idx, clean_name)

        kids_ru = ['MULT', 'DETSK', 'NICK', 'DISNEY', 'ANI', 'TIJI', 'GULLI', 'BOOMERANG', 'CARTOON', 'TLUM', 'SKAZKI']
        for idx, ki in enumerate(kids_ru):
            if ki in ru_k:
                return (5, idx, clean_name)

        music_ru = ['1HD', 'MUZ', 'EUROPAPLUS', 'BRIDGE', 'SHANSON', 'MTV', 'FIRSTMUSIC', 'MUSICBOX', 'MUSIC', 'RUTV']
        for idx, m in enumerate(music_ru):
            if m in ru_k:
                return (6, idx, clean_name)

        return (7, 0, clean_name)

    elif 'SPORT' in category_key:
        if 'BUNDESLIGA' in u:
            if u == 'SKY SPORT BUNDESLIGA': return (1, 0, clean_name)
            m = re.search(r'BUNDESLIGA\s*(\d+)', u)
            if m: return (1, int(m.group(1)), clean_name)
            return (1, 25, clean_name)
            
        sky_sports = [
            'SKY SPORT TOP EVENT', 'SKY SPORT PREMIER LEAGUE', 'SKY SPORT F1',
            'SKY SPORT TENNIS', 'SKY SPORT GOLF', 'SKY SPORT MIX', 'SKY SPORT NEWS',
            'SKY SPORT 1', 'SKY SPORT 2', 'SKY SPORT 3', 'SKY SPORT 4', 'SKY SPORT 5',
            'SKY SPORT 6', 'SKY SPORT 7', 'SKY SPORT 8', 'SKY SPORT 9', 'SKY SPORT 10'
        ]
        for idx, name in enumerate(sky_sports):
            if u == name or u.startswith(name):
                return (2, idx, clean_name)

        dazn = ['DAZN 1', 'DAZN 2', 'DAZN BAR 1', 'DAZN BAR 2', 'DAZN EXCLUSIVE', 'SPORTDIGITAL FUSSBALL']
        for idx, name in enumerate(dazn):
            if name in u:
                return (3, idx, clean_name)

        austria = ['SKY SPORT AUSTRIA', 'BLUE SPORT', 'MYSPORTS']
        for idx, name in enumerate(austria):
            if name in u:
                return (4, idx, clean_name)

        free_sport = ['EUROSPORT 1', 'EUROSPORT 2', 'SPORT1', 'NFL NETWORK', 'REDBULLTV']
        for idx, name in enumerate(free_sport):
            if name in u:
                return (5, idx, clean_name)

        return (6, 0, clean_name)

    elif 'SKY' in category_key:
        sky_cinema = [
            'SKY CINEMA PREMIERE', 'SKY CINEMA PREMIEREN', 'SKY CINEMA HIGHLIGHTS', 'SKY CINEMA ACTION',
            'SKY CINEMA FAMILY', 'SKY CINEMA CLASSICS', 'SKY CINEMA BLOCKBUSTER', 'SKY CINEMA BEST OF',
            'SKY CINEMA SPECIAL', 'SKY CINEMA THRILLER', 'SKY CINEMA FEELGOOD', 'SKY ATLANTIC',
            'SKY ONE', 'SKY REPLAY', 'SKY SHOWCASE', 'SKY KRIMI', 'WARNER TV SERIE',
            'WARNER TV FILM', 'WARNER TV COMEDY', 'UNIVERSAL TV', '13TH STREET', 'SYFY', 'KINOWELT'
        ]
        for idx, name in enumerate(sky_cinema):
            norm_c = re.sub(r'[^A-Z0-9\+]', '', name.upper())
            if norm_c in norm_k or norm_k in norm_c:
                return (1, idx, clean_name)
        return (2, 0, clean_name)

    elif 'DOKU' in category_key:
        doku_lcn = [
            'ZDFINFO', 'KABELEINSDOKU', 'N24DOKU', 'WELTDERWUNDER',
            'DISCOVERYCHANNEL', 'NATIONALGEOGRAPHIC', 'NATGEOWILD', 'ANIMALPLANET', 'HISTORYCHANNEL',
            'SKYDOCUMENTARIES', 'SKYNATURE', 'SKYCRIME',
            'GEOTELEVISION', 'SPIEGELGESCHICHTE', 'SPIEGELTVKONFLIKTE',
            'CURIOSITYCHANNEL', 'CURIOSITYNOW',
            'MARCOPOLOTV', 'TRAVELXP4K', 'BERGBLICK', 'TERRAMATERWILD',
            'BBCHISTORY', 'BBCTRAVEL', 'MAGELLANTVNOW', 'ONETERRA', 'XPLORE',
            'LOVETHEPLANET', 'INSIGHTTV', 'WAIDWERK',
            'CRIMEINVESTIGATION', 'CRIMESCENETV', 'TOPTRUECRIME',
            'AMERICANCRIMES', 'FILMRISETRUECRIME', 'TAETERJAGDCRIMESCENESOLVERS', 'TTERJAGDCRIMESCENESOLVERS',
            'FOCUSTVREPORTAGE', 'WEDOTVBIGSTORIES'
        ]
        if norm_k in doku_lcn:
            return (1, doku_lcn.index(norm_k), clean_name)
        return (2, 0, clean_name)

    elif 'KIDS' in category_key:
        kids_top = [
            'KIKA', 'DISNEY CHANNEL', 'RTL SUPER', 'TOGGO PLUS', 'NICKELODEON',
            'NICK JR', 'CARTOON NETWORK', 'BOOMERANG', 'JUNIOR', 'RIC', 'FIX & FOXI'
        ]
        for idx, name in enumerate(kids_top):
            if name in u:
                return (1, idx, clean_name)
        return (2, 0, clean_name)

    elif 'MUSIK' in category_key:
        top_german_music = [
            'MTV', 'DELUXE MUSIC', 'JUKE BOX', 'GOLDSTAR TV', 'GUTE LAUNE TV', 'GUTE LAUNE', 'STINGRAY CLASSICA'
        ]
        for idx, name in enumerate(top_german_music):
            norm_gm = re.sub(r'[^A-Z0-9\+]', '', name.upper())
            if norm_gm in norm_k or norm_k == norm_gm:
                return (1, idx, clean_name)

        top_uk_music = [
            'MTV 4K', 'MTV 90S', 'MTV BASE', 'MTV DANCE', 'MTV HITS', 'MTV LIVE 4K', 'MTV MUSIC',
            'NOW 70S', 'NOW 80S', 'NOW 90S', '4 MUSIC', 'CLUBLAND TV', 'TRACE URBAN', 'VH1',
            'THAT\'S 60S', 'THAT\'S 80S', 'AYOZAT TV', 'CMT MUSIC', 'RMTV'
        ]
        for idx, name in enumerate(top_uk_music):
            norm_uk = re.sub(r'[^A-Z0-9\+]', '', name.upper())
            if norm_uk in norm_k:
                return (2, idx, clean_name)

        return (3, 0, clean_name)

    return (1, 0, clean_name)

def sanitize_sheet_name(name):
    clean = re.sub(r'[\\/\?\*\:\[\]]', '-', name)
    return clean[:31].strip()

def main():
    print("1. Lade bestehende Excel-Datei & bereinige ausgemusterte Kategorien...")
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws_cat_map = wb["Live TV Kategorien"]

    # Bestehende Kategorien auslesen & Blacklist filtern
    existing_rows = []
    known_cids = set()

    for row in ws_cat_map.iter_rows(min_row=2, values_only=True):
        cname, cid, lang, count, main_cat, sort_order = row[:6]
        cid_str = str(cid).strip()
        main_cat_str = str(main_cat).strip() if main_cat else ''
        
        # Blacklist & User-Wunsch: Spotify und alle 'kann entfernt werden' rigoros löschen!
        if cid_str in BLACKLIST_CATEGORY_IDS or 'kann entfernt werden' in main_cat_str.lower() or 'SPOTIFY' in str(cname).upper():
            continue

        known_cids.add(cid_str)
        existing_rows.append({
            'name': cname,
            'id': cid_str,
            'lang': lang,
            'count': count,
            'main': main_cat_str,
            'sort': sort_order
        })

    print("2. Prüfe Xtream API auf NEUE deutsch- oder russischsprachige Kategorien...")
    cats_url = f"{SERVER_URL}/player_api.php?username={USERNAME}&password={PASSWORD}&action=get_live_categories"
    req_c = urllib.request.Request(cats_url, headers=HEADERS)
    with urllib.request.urlopen(req_c, timeout=30) as resp:
        api_cats = json.loads(resp.read().decode('utf-8'))

    new_cats_found = 0
    for c in api_cats:
        cid = str(c.get('category_id')).strip()
        cname = c.get('category_name', '').strip()
        if cid in known_cids or cid in BLACKLIST_CATEGORY_IDS:
            continue

        main_cat, sort_order = auto_classify_category(cname)
        if main_cat:
            new_cats_found += 1
            print(f"  [NEUE KATEGORIE ENTDECKT] CID {cid}: '{cname}' -> Automatisch zugeordnet zu: '{main_cat}' (Sort {sort_order})")
            existing_rows.append({
                'name': cname,
                'id': cid,
                'lang': 'DE' if main_cat != 'Russian' else 'RU',
                'count': 0,
                'main': main_cat,
                'sort': sort_order
            })
            known_cids.add(cid)

    # Schreibe bereinigten Reiter 'Live TV Kategorien' komplett sauber neu
    ws_cat_map.delete_rows(2, ws_cat_map.max_row)
    for r in existing_rows:
        ws_cat_map.append([r['name'], r['id'], r['lang'], r['count'], r['main'], r['sort']])

    # Kategorien nach Hauptkategorie bündeln
    categories_by_main = {}
    subcat_map = {}
    for r in existing_rows:
        main_cat = r['main']
        sort_val = 999 if r['sort'] is None else int(r['sort'])
        key = (sort_val, main_cat)
        categories_by_main.setdefault(key, []).append({
            'name': r['name'],
            'id': r['id'],
            'count': r['count']
        })
        subcat_map[r['id']] = (sort_val, main_cat, r['name'])

    sorted_groups = sorted(categories_by_main.keys(), key=lambda x: (x[0], x[1]))

    print(f"3. Lade aktuelle Live-Streams von Xtream API...")
    streams_url = f"{SERVER_URL}/player_api.php?username={USERNAME}&password={PASSWORD}&action=get_live_streams"
    req_s = urllib.request.Request(streams_url, headers=HEADERS)
    with urllib.request.urlopen(req_s, timeout=60) as resp:
        all_streams = json.loads(resp.read().decode('utf-8'))
    print(f"-> {len(all_streams)} Live-Streams empfangen.")

    # Stream-Routing durchführen
    items_by_main = {main_name: [] for (sort_val, main_name) in sorted_groups}

    for s in all_streams:
        cid = str(s.get('category_id')).strip()
        if cid not in subcat_map or cid in BLACKLIST_CATEGORY_IDS:
            continue
        sort_val, default_main, subcat_name = subcat_map[cid]

        orig_name = s.get('name', '').strip()
        # Spotify-Streams komplett überspringen
        if 'SPOTIFY' in orig_name.upper():
            continue

        clean_name = clean_channel_name(orig_name, default_main=default_main)
        if clean_name.startswith('---'):
            continue
        norm_k = normalize_key(clean_name)
        score, stype = evaluate_source(orig_name, subcat_name)

        target_mains = [default_main]

        if default_main == 'Russian':
            target_mains = ['Russian']
        elif norm_k in DOKU_CHANNELS:
            target_mains = ['Doku']
            clean_name = DOKU_CHANNELS[norm_k]
            norm_k = normalize_key(clean_name)
        elif norm_k in SPECIAL_REROUTES:
            target_mains = [SPECIAL_REROUTES[norm_k]]
        elif cid == '509': # Sky HEVC Aufteilung
            if norm_k in GERMAN_MUSIC_KEYS:
                target_mains = ['FreeTV', 'Musik']
            elif norm_k in FREETV_ROUTING_KEYS:
                target_mains = ['FreeTV']
            elif norm_k in KIDS_ROUTING_KEYS:
                target_mains = ['Kids']
            elif norm_k in SPORT_ROUTING_KEYS:
                target_mains = ['Sport']
            else:
                target_mains = ['Sky']
        elif default_main != 'Russian' and norm_k in GERMAN_MUSIC_KEYS:
            if 'Musik' in items_by_main and 'Musik' not in target_mains:
                target_mains.append('Musik')

        for tm in target_mains:
            if tm in items_by_main:
                items_by_main[tm].append({
                    'orig': orig_name,
                    'clean': clean_name,
                    'norm_k': norm_k,
                    'subcat': subcat_name,
                    'stream_id': s.get('stream_id', ''),
                    'stream_icon': s.get('stream_icon', ''),
                    'epg_channel_id': s.get('epg_channel_id', ''),
                    'score': score,
                    'stype': stype
                })

    # Styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B71C1C", end_color="B71C1C", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    font_main = Font(name="Calibri", size=10, bold=True, color="1B5E20")
    font_backup = Font(name="Calibri", size=10, color="5D4037")
    font_single = Font(name="Calibri", size=10, bold=True, color="212121")
    font_regular = Font(name="Calibri", size=10)
    font_muted = Font(name="Calibri", size=9, italic=True, color="757575")

    fill_main = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    fill_backup = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")
    fill_single = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_alt = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")

    for s in [name for name in wb.sheetnames if name != "Live TV Kategorien"]:
        del wb[s]

    headers = [
        "Bereinigter Sendername",
        "Multi-Stream Rolle",
        "Erkannte Quelle / Profil",
        "Unterkategorie",
        "Originaler Sendername",
        "Stream_ID"
    ]

    summary_stats = []
    sheet_number = 1

    # M3U Listen initialisieren
    m3u_clean_lines = ["#EXTM3U\n"]
    m3u_multi_lines = ["#EXTM3U\n"]

    print("4. Erstelle sortierte Blätter und generiere M3U-Playlists...")
    for sort_val, main_name in sorted_groups:
        sheet_title = sanitize_sheet_name(f"{sheet_number}. {main_name}")
        sheet_number += 1

        ws = wb.create_sheet(title=sheet_title)
        ws.sheet_properties.tabColor = "B71C1C"
        ws.freeze_panes = "A2"

        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for col_idx in range(1, 7):
            c = ws.cell(row=1, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_alignment
            c.border = thin_border

        raw_items = items_by_main.get(main_name, [])

        channel_groups = {}
        for item in raw_items:
            if main_name == 'Russian':
                k = item['clean']
            else:
                k = item['norm_k'] or item['clean']
            channel_groups.setdefault(k, []).append(item)

        total_unique_channels = len(channel_groups)
        channels_with_backups = sum(1 for v in channel_groups.values() if len(v) > 1)
        total_backup_streams = sum(len(v) - 1 for v in channel_groups.values() if len(v) > 1)

        summary_stats.append({
            'title': sheet_title,
            'total_streams': len(raw_items),
            'unique_channels': total_unique_channels,
            'channels_with_backups': channels_with_backups,
            'backup_streams': total_backup_streams
        })

        cat_upper = main_name.upper()
        sorted_keys = sorted(
            channel_groups.keys(),
            key=lambda k: get_channel_sort_priority(channel_groups[k][0]['clean'], cat_upper)
        )

        current_row = 2
        for k in sorted_keys:
            group = channel_groups[k]
            group.sort(key=lambda x: x['score'], reverse=True)
            has_duplicates = len(group) > 1

            # M3U CLEAN: Nur die beste Quelle 1 hinzufügen!
            best_item = group[0]
            stream_url = f"{SERVER_URL}/live/{USERNAME}/{PASSWORD}/{best_item['stream_id']}.ts"
            epg_id = best_item['epg_channel_id'] or ""
            logo = best_item['stream_icon'] or ""
            clean_display_name = best_item['clean']
            
            m3u_clean_lines.append(
                f'#EXTINF:-1 tvg-id="{epg_id}" tvg-name="{clean_display_name}" tvg-logo="{logo}" group-title="{sheet_title}",{clean_display_name}\n'
            )
            m3u_clean_lines.append(f"{stream_url}\n")

            for idx, item in enumerate(group):
                ws.row_dimensions[current_row].height = 20

                if not has_duplicates:
                    role_text = "⚪ Einzelsender"
                    row_fill = fill_single
                    name_font = font_single
                    m3u_name = item['clean']
                elif idx == 0:
                    role_text = f"🟢 Hauptsender (Quelle 1 - Beste von {len(group)})"
                    row_fill = fill_main
                    name_font = font_main
                    m3u_name = f"{item['clean']} [{item['stype'].split('(')[0].strip()}]"
                else:
                    role_text = f"🟡 Backup (Quelle {idx + 1})"
                    row_fill = fill_backup
                    name_font = font_backup
                    m3u_name = f"   ↳ {item['clean']} (Backup {idx+1}: {item['stype'].split('(')[0].strip()})"

                # M3U MULTISTREAM: Alle Streams hinzufügen
                item_url = f"{SERVER_URL}/live/{USERNAME}/{PASSWORD}/{item['stream_id']}.ts"
                m3u_multi_lines.append(
                    f'#EXTINF:-1 tvg-id="{item["epg_channel_id"] or ""}" tvg-name="{item["clean"]}" tvg-logo="{item["stream_icon"] or ""}" group-title="{sheet_title}",{m3u_name}\n'
                )
                m3u_multi_lines.append(f"{item_url}\n")

                row_vals = [
                    item['clean'] if idx == 0 else f"   ↳ {item['clean']}",
                    role_text,
                    item['stype'],
                    item['subcat'],
                    item['orig'],
                    str(item['stream_id'])
                ]
                ws.append(row_vals)

                c_name = ws.cell(row=current_row, column=1)
                c_role = ws.cell(row=current_row, column=2)
                c_stype = ws.cell(row=current_row, column=3)
                c_subcat = ws.cell(row=current_row, column=4)
                c_orig = ws.cell(row=current_row, column=5)
                c_sid = ws.cell(row=current_row, column=6)

                for cell in [c_name, c_role, c_stype, c_subcat, c_orig, c_sid]:
                    cell.border = thin_border
                    cell.fill = row_fill

                c_name.font = name_font
                c_role.font = font_regular
                c_stype.font = font_regular
                c_subcat.font = font_muted
                c_orig.font = font_regular
                c_sid.font = font_muted
                c_sid.alignment = Alignment(horizontal="center", vertical="center")

                current_row += 1

        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions["E"].width = 44
        ws.column_dimensions["F"].width = 14

        print(f"-> Blatt '{sheet_title}': {len(raw_items)} Streams -> {total_unique_channels} einzigartige Sender")

    # Übersichtstabelle
    print("5. Erstelle Übersichtsblatt '📊 Multi-Stream Übersicht'...")
    ws_ov = wb.create_sheet(title="📊 Multi-Stream Übersicht", index=0)
    ws_ov.sheet_properties.tabColor = "1B5E20"
    ws_ov.freeze_panes = "A2"

    ov_headers = [
        "Überkategorie",
        "Gesamte Streams (roh)",
        "Echte Einzigartige Sender",
        "Sender mit Backup-Quellen",
        "Gesamte Backup-Streams",
        "Duplikat-Quote"
    ]
    ws_ov.append(ov_headers)
    ws_ov.row_dimensions[1].height = 28

    for col_idx in range(1, 7):
        c = ws_ov.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
        c.alignment = header_alignment
        c.border = thin_border

    total_all_streams = 0
    total_all_unique = 0
    total_all_with_backups = 0
    total_all_backups = 0

    for idx, stat in enumerate(summary_stats, 2):
        ws_ov.row_dimensions[idx].height = 22
        pct = (stat['backup_streams'] / stat['total_streams'] * 100) if stat['total_streams'] > 0 else 0
        
        total_all_streams += stat['total_streams']
        total_all_unique += stat['unique_channels']
        total_all_with_backups += stat['channels_with_backups']
        total_all_backups += stat['backup_streams']

        ws_ov.append([
            stat['title'],
            stat['total_streams'],
            stat['unique_channels'],
            stat['channels_with_backups'],
            stat['backup_streams'],
            f"{pct:.1f}%"
        ])

        r_fill = fill_alt if idx % 2 == 0 else fill_single
        for col_idx in range(1, 7):
            cell = ws_ov.cell(row=idx, column=col_idx)
            cell.border = thin_border
            cell.fill = r_fill
            cell.font = font_regular
            if col_idx in [2, 3, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    total_row_idx = len(summary_stats) + 2
    ws_ov.row_dimensions[total_row_idx].height = 24
    total_pct = (total_all_backups / total_all_streams * 100) if total_all_streams > 0 else 0

    ws_ov.append([
        "GESAMT",
        total_all_streams,
        total_all_unique,
        total_all_with_backups,
        total_all_backups,
        f"{total_pct:.1f}%"
    ])
    bold_total_font = Font(name="Calibri", size=11, bold=True)
    total_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    for col_idx in range(1, 7):
        cell = ws_ov.cell(row=total_row_idx, column=col_idx)
        cell.border = thin_border
        cell.fill = total_fill
        cell.font = bold_total_font
        if col_idx in [2, 3, 4, 5, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_ov.column_dimensions["A"].width = 28
    ws_ov.column_dimensions["B"].width = 24
    ws_ov.column_dimensions["C"].width = 26
    ws_ov.column_dimensions["D"].width = 26
    ws_ov.column_dimensions["E"].width = 24
    ws_ov.column_dimensions["F"].width = 18

    print("6. Speichere Excel-Arbeitsmappe & exportiere M3U8-Dateien...")
    try:
        wb.save(XLSX_PATH)
        saved_xlsx = XLSX_PATH
    except PermissionError:
        saved_xlsx = os.path.join(PROJECT_ROOT, "live_tv_categories_export_updated.xlsx")
        wb.save(saved_xlsx)
        print(f"\nHINWEIS: '{os.path.basename(XLSX_PATH)}' ist aktuell in Excel geöffnet.")
        print(f"-> Aktuelle Version wurde gespeichert unter: '{os.path.basename(saved_xlsx)}'")
    
    with open(M3U_CLEAN_PATH, 'w', encoding='utf-8') as f:
        f.writelines(m3u_clean_lines)

    with open(M3U_MULTI_PATH, 'w', encoding='utf-8') as f:
        f.writelines(m3u_multi_lines)

    print(f"\nSYNCHRONISATION ERFOLGREICH ABGESCHLOSSEN!")
    print(f"Excel-Datei:        {saved_xlsx}")
    print(f"M3U Clean Playlist: {M3U_CLEAN_PATH} ({len(m3u_clean_lines)//2} Kanäle)")
    print(f"M3U Multi Playlist: {M3U_MULTI_PATH} ({len(m3u_multi_lines)//2} Streams)")

if __name__ == "__main__":
    main()
